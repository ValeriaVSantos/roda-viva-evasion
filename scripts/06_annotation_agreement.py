"""Analyze three human annotation columns in the project workbook.

The public export excludes the question and response text. It contains only
source identifiers, anonymized labels, and provisional consensus information.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

LABELS = ("RESPONDE", "PARCIAL", "ESQUIVA")
LABEL_ALIASES = {
    "RESPONDE": "RESPONDE",
    "RESPONDS": "RESPONDE",
    "PARCIAL": "PARCIAL",
    "PARTIAL": "PARCIAL",
    "ESQUIVA": "ESQUIVA",
    "EVADES": "ESQUIVA",
}


def normalize_label(value):
    if value is None or not str(value).strip():
        return None
    normalized = str(value).strip().upper()
    if normalized not in LABEL_ALIASES:
        raise ValueError(f"unknown label: {value!r}")
    return LABEL_ALIASES[normalized]


def cohen_kappa(left, right):
    pairs = [(a, b) for a, b in zip(left, right) if a is not None and b is not None]
    if not pairs:
        return None, None, 0
    observed = sum(a == b for a, b in pairs) / len(pairs)
    left_counts, right_counts = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    expected = sum((left_counts[label] / len(pairs)) * (right_counts[label] / len(pairs)) for label in LABELS)
    kappa = (observed - expected) / (1 - expected) if expected < 1 else 1.0
    return kappa, observed, len(pairs)


def fleiss_kappa(rows):
    complete = [ratings for ratings in rows if len(ratings) == 3]
    if not complete:
        return None, 0
    per_item = []
    totals = Counter()
    for ratings in complete:
        counts = Counter(ratings)
        totals.update(ratings)
        per_item.append((sum(value * value for value in counts.values()) - 3) / 6)
    observed = sum(per_item) / len(per_item)
    expected = sum((totals[label] / (len(complete) * 3)) ** 2 for label in LABELS)
    value = (observed - expected) / (1 - expected) if expected < 1 else 1.0
    return value, len(complete)


def krippendorff_alpha_nominal(rows):
    usable = [ratings for ratings in rows if len(ratings) >= 2]
    coincidence = {a: {b: 0.0 for b in LABELS} for a in LABELS}
    for ratings in usable:
        denominator = len(ratings) - 1
        for index, left in enumerate(ratings):
            for other_index, right in enumerate(ratings):
                if index != other_index:
                    coincidence[left][right] += 1 / denominator
    marginals = {label: sum(coincidence[label].values()) for label in LABELS}
    total = sum(marginals.values())
    if total <= 1:
        return None, len(usable)
    observed_disagreement = sum(coincidence[a][b] for a in LABELS for b in LABELS if a != b) / total
    expected_disagreement = sum(marginals[a] * marginals[b] / (total - 1) for a in LABELS for b in LABELS if a != b) / total
    alpha = 1 - observed_disagreement / expected_disagreement if expected_disagreement else 1.0
    return alpha, len(usable)


def provisional_consensus(ratings):
    if len(ratings) < 2:
        return None, "INSUFFICIENT"
    counts = Counter(ratings)
    label, count = counts.most_common(1)[0]
    if len(ratings) == 2 and count == 2:
        return label, "AGREEMENT_2_OF_2"
    if len(ratings) == 3 and count == 3:
        return label, "UNANIMOUS_3_OF_3"
    if len(ratings) == 3 and count == 2:
        return label, "MAJORITY_2_OF_3"
    return None, "ADJUDICATION_REQUIRED"


def load_rows(path: Path, sheet_name: str):
    sheet = load_workbook(path, read_only=True, data_only=True)[sheet_name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    expected = ["ENTREVISTADO", "DATA", "ORDEM_PERGUNTA", "ORDEM_RESPOSTA", "ANOTACAO_HUMANA", "Anotacao 2", "Anotador 3"]
    missing = [column for column in expected if column not in headers]
    if missing:
        raise ValueError(f"missing workbook columns: {missing}")
    positions = {name: headers.index(name) for name in expected}
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row[index] is not None for index in positions.values()):
            continue
        ratings = [normalize_label(row[positions[name]]) for name in ("ANOTACAO_HUMANA", "Anotacao 2", "Anotador 3")]
        records.append({
            "interviewee": row[positions["ENTREVISTADO"]],
            "date": row[positions["DATA"]],
            "question_order": row[positions["ORDEM_PERGUNTA"]],
            "response_order": row[positions["ORDEM_RESPOSTA"]],
            "ratings": ratings,
        })
    return records


def analyze(records):
    rating_rows = [[label for label in record["ratings"] if label] for record in records]
    coverage = Counter(len(row) for row in rating_rows)
    pairwise = {}
    for left, right in combinations(range(3), 2):
        kappa, agreement, count = cohen_kappa([r["ratings"][left] for r in records], [r["ratings"][right] for r in records])
        pairwise[f"annotator_{left + 1}_vs_{right + 1}"] = {"n": count, "observed_agreement": agreement, "cohen_kappa": kappa}
    alpha, alpha_n = krippendorff_alpha_nominal(rating_rows)
    fleiss, fleiss_n = fleiss_kappa(rating_rows)
    statuses = Counter(provisional_consensus(row)[1] for row in rating_rows)
    distributions = {f"annotator_{index + 1}": Counter(record["ratings"][index] for record in records if record["ratings"][index]) for index in range(3)}
    return {
        "items_total": len(records),
        "coverage": {"zero_ratings": coverage[0], "one_rating": coverage[1], "two_ratings": coverage[2], "three_ratings": coverage[3]},
        "label_distribution": distributions,
        "agreement": {"krippendorff_alpha_nominal": alpha, "krippendorff_items": alpha_n, "fleiss_kappa_complete_items": fleiss, "fleiss_items": fleiss_n, "pairwise": pairwise},
        "provisional_consensus": statuses,
    }


def export_public_labels(records, output_path: Path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fields = ["item_id", "interviewee", "date", "question_order", "response_order", "annotator_a", "annotator_b", "annotator_c", "provisional_label", "consensus_status"]
    with output_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for index, record in enumerate(records, start=1):
            present = [label for label in record["ratings"] if label]
            consensus, status = provisional_consensus(present)
            writer.writerow({
                "item_id": f"RV-{index:04d}",
                "interviewee": record["interviewee"],
                "date": record["date"],
                "question_order": record["question_order"],
                "response_order": record["response_order"],
                "annotator_a": record["ratings"][0],
                "annotator_b": record["ratings"][1],
                "annotator_c": record["ratings"][2],
                "provisional_label": consensus,
                "consensus_status": status,
            })


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--sheet", default="✏️ Anotação")
    parser.add_argument("--output-data", type=Path, required=True)
    parser.add_argument("--output-summary", type=Path, required=True)
    return parser.parse_args()


def main():
    args = parse_args()
    records = load_rows(args.input, args.sheet)
    summary = analyze(records)
    export_public_labels(records, args.output_data)
    args.output_summary.parent.mkdir(parents=True, exist_ok=True)
    args.output_summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=dict) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=dict))


if __name__ == "__main__":
    main()
